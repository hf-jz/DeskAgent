#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Minimal PPTX generator for the creation workspace (python-pptx in the
dedicated venv: ~/.hermes/pptx-venv). Usage: pptx-gen.py <blocks.json> <out.pptx>"""
import json, sys
from pptx import Presentation
from pptx.util import Inches, Pt

def main():
    data = json.load(open(sys.argv[1]))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3), Inches(2))
    p = tb.text_frame.paragraphs[0]
    p.text = data.get('title', '作品')
    p.font.size = Pt(44); p.font.bold = True
    for b in data.get('blocks', []):
        s = prs.slides.add_slide(blank)
        if b.get('kind') == 'image':
            try:
                s.shapes.add_picture(b['content'], Inches(2), Inches(1), width=Inches(9))
            except Exception:
                tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(1))
                tb.text_frame.paragraphs[0].text = '图片无法嵌入: ' + b['content']
        else:
            tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(11.3), Inches(5.5))
            lines = str(b.get('content', '')).split('\n')
            p = tb.text_frame.paragraphs[0]
            p.text = lines[0][:40] if lines and lines[0] else '内容'
            p.font.size = Pt(32); p.font.bold = True
            for line in lines[1:]:
                pp = tb.text_frame.add_paragraph(); pp.text = line; pp.font.size = Pt(18)
    prs.save(sys.argv[2])
    print(sys.argv[2])

if __name__ == '__main__':
    main()
